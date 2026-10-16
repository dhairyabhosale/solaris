"""Retrieve and parse the Grid-India / NERLDC SCADA workbooks.

Source
------
Mukherjee, D.; Kalita, K.; Kumar, S. (2024), "Electricity Demand, Solar and
Wind Generation Data (2021-2023) of India at 1-hour interval", Mendeley Data,
V1, doi: 10.17632/y58jknpgs8.1, licensed CC BY 4.0. The dataset description
credits the North-Eastern Regional Load Despatch Centre (NERLDC), Grid-India,
Ministry of Power, Government of India as the originating source.

Verified discrepancies in that dataset (see docs/DATA_SOURCES.md)
----------------------------------------------------------------
1. The title says "1-hour interval". The files are NOT hourly. Each workbook
   embeds an `_osi_config` sheet carrying the OSIsoft export sample period in
   seconds: `"r":300.0` (5-minute) for the earlier months and `"r":10.0`
   (10-second) for the later ones. This module reads that field, records it,
   and aggregates to a genuine hourly series itself.
2. The sampling interval therefore *changes mid-dataset*. Hourly aggregation
   makes the two regimes comparable, but the number of raw samples backing each
   hourly value differs by a factor of 30, which is recorded per hour as
   `n_samples` so that downstream code can audit it.
3. Columns H, I and J of Sheet1 ("Solar", "Wind+Solar"/"Solar+Wind", "Total")
   are Excel formulas, not stored values. They are recomputed here from the raw
   SCADA columns rather than trusted, so the derivation is explicit.
4. The raw solar point `ALL_IND_SOLAR|P` goes slightly negative at night
   (around -10 to -25 MW), which is a SCADA measurement offset rather than real
   generation. Both the raw and the clipped-at-zero series are kept.
"""

from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass, asdict
from pathlib import Path

import numpy as np
import pandas as pd

from solaris import config

logger = logging.getLogger(__name__)

# "September 2021.xlsx", and the dataset's own typo "Novemeber 2023.xlsx".
_MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10,
    "november": 11, "novemeber": 11,  # sic -- the dataset misspells November 2023
    "december": 12,
}


def parse_month_from_filename(filename: str) -> tuple[int, int]:
    """Return ``(year, month)`` parsed from a Mendeley workbook filename.

    Tolerates the dataset's "Novemeber 2023.xlsx" misspelling.
    """
    stem = Path(filename).stem.strip().lower()
    match = re.match(r"([a-z]+)\s+(\d{4})$", stem)
    if not match:
        raise ValueError(f"Cannot parse month/year from filename: {filename!r}")
    name, year = match.group(1), int(match.group(2))
    if name not in _MONTHS:
        raise ValueError(f"Unrecognised month name {name!r} in {filename!r}")
    return year, _MONTHS[name]


# --------------------------------------------------------------------------
# Download
# --------------------------------------------------------------------------


# Mendeley's edge sits behind a bot filter that rejects `requests` outright
# with a 403 regardless of User-Agent -- it appears to fingerprint the TLS
# handshake rather than the header. `httpx` negotiates acceptably and is already
# a dependency of the web app, so the pipeline uses it for every Mendeley call.
_USER_AGENT = (
    "solaris-grid-intelligence/1.0 (research pipeline; "
    "+https://github.com/dhairyabhosale/solaris)"
)
_HTTP_HEADERS = {"User-Agent": _USER_AGENT, "Accept": "application/json"}


def list_remote_files() -> list[dict]:
    """Return the Mendeley file manifest (filename, size, sha256, download URL)."""
    import httpx

    response = httpx.get(
        config.MENDELEY_FILES_API, timeout=120, headers=_HTTP_HEADERS, follow_redirects=True
    )
    response.raise_for_status()
    return response.json()


def download_all(dest_dir: Path | None = None, *, force: bool = False) -> list[Path]:
    """Download every monthly workbook, skipping files already present and intact.

    The raw workbooks total roughly 750 MB and are deliberately NOT tracked in
    git; this function is the reproducible way to recreate ``data/raw``.
    """
    import httpx

    dest_dir = dest_dir or config.GRID_RAW_DIR
    dest_dir.mkdir(parents=True, exist_ok=True)

    manifest = list_remote_files()
    manifest_path = dest_dir / "_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    paths: list[Path] = []
    for entry in sorted(manifest, key=lambda e: e["filename"]):
        name = entry["filename"]
        expected_size = int(entry["size"])
        target = dest_dir / name
        if target.exists() and target.stat().st_size == expected_size and not force:
            logger.info("skip %s (already downloaded)", name)
            paths.append(target)
            continue

        url = entry["content_details"]["download_url"]
        logger.info("download %s (%.1f MB)", name, expected_size / 1e6)
        tmp = target.with_suffix(target.suffix + ".part")
        with httpx.stream(
            "GET",
            url,
            timeout=900,
            headers={"User-Agent": _USER_AGENT},
            follow_redirects=True,
        ) as response:
            response.raise_for_status()
            with tmp.open("wb") as fh:
                for chunk in response.iter_bytes(chunk_size=1 << 20):
                    fh.write(chunk)
        tmp.replace(target)
        paths.append(target)

    return paths


# --------------------------------------------------------------------------
# Parse
# --------------------------------------------------------------------------


@dataclass
class WorkbookInfo:
    """What a single monthly workbook turned out to contain."""

    filename: str
    year: int
    month: int
    sample_period_seconds: float | None
    raw_rows: int
    hours: int
    first_timestamp: str
    last_timestamp: str


def read_sample_period(path: Path) -> float | None:
    """Read the OSIsoft export sample period (seconds) from ``_osi_config``.

    Returns ``None`` when the sheet or the ``r`` key is absent, rather than
    guessing an interval.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        if "_osi_config" not in workbook.sheetnames:
            return None
        sheet = workbook["_osi_config"]
        for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
            for cell in row:
                if isinstance(cell, str) and cell.lstrip().startswith("{"):
                    try:
                        return float(json.loads(cell).get("r"))
                    except (ValueError, TypeError, json.JSONDecodeError):
                        continue
        return None
    finally:
        workbook.close()


def _iter_sheet1_records(path: Path):
    """Yield ``(timestamp, [values...])`` from Sheet1, skipping the two headers.

    Streams with ``read_only=True`` because the 10-second workbooks hold ~268k
    rows each and must not be materialised in full.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["Sheet1"]
        wanted = config.SCADA_COLUMN_INDEX
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i < 2:  # two header rows
                continue
            timestamp = row[0]
            if timestamp is None:
                continue
            values = []
            for _, idx in wanted.items():
                value = row[idx] if idx < len(row) else None
                # Formula strings and SCADA "#N/A" placeholders become NaN.
                values.append(value if isinstance(value, (int, float)) else np.nan)
            yield timestamp, values
    finally:
        workbook.close()


def parse_workbook(path: Path) -> tuple[pd.DataFrame, WorkbookInfo]:
    """Parse one monthly workbook into an hourly-mean frame, IST-localised.

    Returns the hourly frame and a record of what the file contained. The frame
    is indexed by tz-aware hour-start timestamps in Asia/Kolkata and carries an
    ``n_samples`` column giving how many raw SCADA samples backed each hour.
    """
    columns = list(config.SCADA_COLUMN_INDEX)
    timestamps: list = []
    rows: list[list[float]] = []
    for timestamp, values in _iter_sheet1_records(path):
        timestamps.append(timestamp)
        rows.append(values)

    if not rows:
        raise ValueError(f"No data rows parsed from {path}")

    raw = pd.DataFrame(rows, columns=columns)
    raw["timestamp"] = pd.to_datetime(pd.Series(timestamps), errors="coerce")
    raw = raw.dropna(subset=["timestamp"])
    raw_rows = len(raw)

    # The workbooks carry naive wall-clock timestamps from despatch-centre
    # systems that run on IST. India has no DST, so localisation is unambiguous.
    raw["timestamp"] = raw["timestamp"].dt.tz_localize(config.IST)

    hourly = (
        raw.set_index("timestamp")
        .resample("1h")
        .agg({c: "mean" for c in columns})
    )
    hourly["n_samples"] = (
        raw.set_index("timestamp")[columns[0]].resample("1h").size()
    )
    hourly = hourly[hourly["n_samples"] > 0]

    year, month = parse_month_from_filename(path.name)
    info = WorkbookInfo(
        filename=path.name,
        year=year,
        month=month,
        sample_period_seconds=read_sample_period(path),
        raw_rows=raw_rows,
        hours=len(hourly),
        first_timestamp=str(hourly.index.min()),
        last_timestamp=str(hourly.index.max()),
    )
    return hourly, info


def derive_columns(hourly: pd.DataFrame) -> pd.DataFrame:
    """Add the clipped-solar, total-generation and residual-load columns.

    These replace the Excel formulas in the source workbook's columns H/I/J so
    the derivation is explicit and testable.
    """
    out = hourly.copy()

    # ALL_IND_SOLAR|P reads slightly negative at night (SCADA offset, not
    # generation). Clip at zero for every downstream use; keep the raw column.
    out["solar_mw"] = out["solar_raw_mw"].clip(lower=0)
    out["vre_mw"] = out["solar_mw"] + out["wind_mw"]
    out["generation_mw"] = (
        out[["thermal_mw", "hydro_mw", "gas_mw", "nuclear_mw", "wind_mw"]].sum(
            axis=1, min_count=5
        )
        + out["solar_mw"]
    )

    # RESEARCH RESIDUAL LOAD. Defined here explicitly as metered all-India
    # demand minus the two variable renewable SCADA aggregates. This is NOT an
    # operator's net-load definition: NLDC_DEMAND|P is itself a met-demand
    # signal, behind-the-meter rooftop solar never appears in either series, and
    # no interchange or storage term is subtracted. It is used only as a
    # consistent indicator of the evening ramp, and is labelled as such.
    out["residual_load_mw"] = out["demand_mw"] - out["vre_mw"]
    return out


def build_hourly(raw_dir: Path | None = None) -> tuple[pd.DataFrame, list[WorkbookInfo]]:
    """Parse every downloaded workbook into one continuous hourly frame."""
    raw_dir = raw_dir or config.GRID_RAW_DIR
    workbooks = sorted(
        p for p in raw_dir.glob("*.xlsx") if not p.name.startswith(("~", "_"))
    )
    if not workbooks:
        raise FileNotFoundError(
            f"No workbooks in {raw_dir}. Run `python -m solaris.cli fetch-grid` first."
        )

    frames: list[pd.DataFrame] = []
    infos: list[WorkbookInfo] = []
    for path in workbooks:
        logger.info("parsing %s", path.name)
        frame, info = parse_workbook(path)
        frames.append(frame)
        infos.append(info)

    combined = pd.concat(frames).sort_index()

    # Months overlap at boundaries only if the source duplicates an hour; keep
    # the first occurrence and report the count in the quality report.
    combined = combined[~combined.index.duplicated(keep="first")]

    # Reindex onto a complete hourly grid so that genuinely missing hours become
    # explicit NaN rows rather than silently vanishing.
    full_index = pd.date_range(
        combined.index.min(), combined.index.max(), freq="1h", tz=config.IST
    )
    combined = combined.reindex(full_index)
    combined["n_samples"] = combined["n_samples"].fillna(0).astype(int)
    combined.index.name = "timestamp_ist"

    combined = derive_columns(combined)
    infos.sort(key=lambda i: (i.year, i.month))
    return combined, infos


def save_hourly(frame: pd.DataFrame, infos: list[WorkbookInfo]) -> Path:
    """Write the hourly grid frame and its per-workbook provenance record."""
    config.ensure_dirs()
    frame.to_parquet(config.GRID_HOURLY_PARQUET)
    (config.INTERIM_DIR / "grid_workbooks.json").write_text(
        json.dumps([asdict(i) for i in infos], indent=2), encoding="utf-8"
    )
    return config.GRID_HOURLY_PARQUET
