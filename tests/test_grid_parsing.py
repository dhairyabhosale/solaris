"""Tests for the Grid-India workbook parser.

These run on small synthetic fixtures, not the 700 MB raw download, so the
suite needs no network access and no data directory.
"""

from __future__ import annotations

import json

import numpy as np
import pandas as pd
import pytest

from solaris import config
from solaris.data import grid


# --------------------------------------------------------------------------
# Filename parsing
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("September 2021.xlsx", (2021, 9)),
        ("December 2023.xlsx", (2023, 12)),
        ("May 2022.xlsx", (2022, 5)),
        # The dataset genuinely ships this misspelling.
        ("Novemeber 2023.xlsx", (2023, 11)),
    ],
)
def test_parse_month_from_filename(filename, expected):
    assert grid.parse_month_from_filename(filename) == expected


def test_parse_month_rejects_unknown_month():
    with pytest.raises(ValueError, match="Unrecognised month"):
        grid.parse_month_from_filename("Smarch 2023.xlsx")


def test_parse_month_rejects_malformed_name():
    with pytest.raises(ValueError, match="Cannot parse"):
        grid.parse_month_from_filename("2023-05.xlsx")


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------


def _make_workbook(tmp_path, name: str, sample_period: float, hours: int = 3):
    """Build a miniature workbook with the real sheet layout and a config blob."""
    import openpyxl

    workbook = openpyxl.Workbook()
    config_sheet = workbook.active
    config_sheet.title = "_osi_config"
    config_sheet["A1"] = 3
    config_sheet["A2"] = json.dumps({"r": sample_period, "tc": []})

    sheet = workbook.create_sheet("Sheet1")
    sheet.append(["Time"] + [f"SCADA/{i}" for i in range(14)])
    sheet.append([
        None, "TOTAL_THM_ONLY|P", "TOTAL_NLDC_HYD|P", "TOTAL_NLDC_GAS|P",
        "TOTAL_NLDC_NPC|P", "ALL_INDIA_WIND|P", "ALL_IND_SOLAR|P",
        "Solar", "Wind+Solar", "Total", "NLDC_DEMAND|P",
        "INTRA_SYSTEM_INTERTIA|P", "ER_Total_Inertia|P", "SYS_IN|P", "SPARE 05",
    ])

    step = pd.Timedelta(seconds=sample_period)
    start = pd.Timestamp("2022-05-01 00:00:00")
    samples_per_hour = int(3600 / sample_period)
    for i in range(hours * samples_per_hour):
        timestamp = (start + i * step).to_pydatetime()
        # Solar is negative at night, as in the real SCADA feed.
        solar = -20.0 if timestamp.hour < 6 else 5000.0
        sheet.append([
            timestamp,
            100000.0,  # thermal
            20000.0,   # hydro
            4000.0,    # gas
            5000.0,    # nuclear
            8000.0,    # wind
            solar,     # raw solar
            "=MAX(0,G3)",      # formula column, must be ignored
            "=SUM(F3,H3)",     # formula column
            "=SUM(B3:F3,H3)",  # formula column
            150000.0,  # demand
            0, 0, 0, 0,
        ])

    path = tmp_path / name
    workbook.save(path)
    return path


# --------------------------------------------------------------------------
# Sample period detection
# --------------------------------------------------------------------------


def test_read_sample_period_five_minute(tmp_path):
    path = _make_workbook(tmp_path, "May 2022.xlsx", 300.0, hours=2)
    assert grid.read_sample_period(path) == 300.0


def test_read_sample_period_ten_second(tmp_path):
    path = _make_workbook(tmp_path, "May 2023.xlsx", 10.0, hours=1)
    assert grid.read_sample_period(path) == 10.0


def test_read_sample_period_missing_config(tmp_path):
    """A workbook without `_osi_config` returns None rather than a guess."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.title = "Sheet1"
    path = tmp_path / "June 2022.xlsx"
    workbook.save(path)
    assert grid.read_sample_period(path) is None


# --------------------------------------------------------------------------
# Workbook parsing
# --------------------------------------------------------------------------


def test_parse_workbook_aggregates_to_hourly(tmp_path):
    path = _make_workbook(tmp_path, "May 2022.xlsx", 300.0, hours=3)
    hourly, info = grid.parse_workbook(path)

    assert len(hourly) == 3
    assert info.sample_period_seconds == 300.0
    assert info.raw_rows == 36  # 3 hours x 12 five-minute samples
    assert (hourly["n_samples"] == 12).all()


def test_parse_workbook_localises_to_ist(tmp_path):
    path = _make_workbook(tmp_path, "May 2022.xlsx", 300.0, hours=2)
    hourly, _ = grid.parse_workbook(path)

    assert hourly.index.tz is not None
    assert str(hourly.index.tz) == config.IST_NAME
    # 00:00 naive IST is 18:30 UTC the previous day.
    assert hourly.index[0].utcoffset() == pd.Timedelta(hours=5, minutes=30)


def test_parse_workbook_ignores_formula_columns(tmp_path):
    """Columns H/I/J hold formula strings; they must not become data."""
    path = _make_workbook(tmp_path, "May 2022.xlsx", 300.0, hours=2)
    hourly, _ = grid.parse_workbook(path)

    # Only the seven real SCADA points plus n_samples should survive.
    assert set(hourly.columns) == set(config.SCADA_COLUMN_INDEX) | {"n_samples"}
    assert hourly["demand_mw"].notna().all()
    assert np.isclose(hourly["demand_mw"].iloc[0], 150000.0)


def test_parse_workbook_handles_ten_second_regime(tmp_path):
    path = _make_workbook(tmp_path, "May 2023.xlsx", 10.0, hours=1)
    hourly, info = grid.parse_workbook(path)

    assert len(hourly) == 1
    assert info.sample_period_seconds == 10.0
    assert hourly["n_samples"].iloc[0] == 360


def test_parse_workbook_empty_raises(tmp_path):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Time"])
    sheet.append([None])
    path = tmp_path / "July 2022.xlsx"
    workbook.save(path)

    with pytest.raises(ValueError, match="No data rows"):
        grid.parse_workbook(path)


# --------------------------------------------------------------------------
# Derived columns
# --------------------------------------------------------------------------


def _hourly_fixture() -> pd.DataFrame:
    index = pd.date_range("2022-05-01", periods=4, freq="1h", tz=config.IST)
    return pd.DataFrame(
        {
            "thermal_mw": [100000.0] * 4,
            "hydro_mw": [20000.0] * 4,
            "gas_mw": [4000.0] * 4,
            "nuclear_mw": [5000.0] * 4,
            "wind_mw": [8000.0] * 4,
            "solar_raw_mw": [-20.0, -15.0, 5000.0, 12000.0],
            "demand_mw": [150000.0, 148000.0, 160000.0, 175000.0],
            "n_samples": [12] * 4,
        },
        index=index,
    )


def test_derive_clips_negative_solar():
    out = grid.derive_columns(_hourly_fixture())
    assert (out["solar_mw"] >= 0).all()
    assert out["solar_mw"].iloc[0] == 0.0
    assert out["solar_mw"].iloc[3] == 12000.0
    # The raw column is preserved for auditing.
    assert out["solar_raw_mw"].iloc[0] == -20.0


def test_derive_vre_uses_clipped_solar():
    out = grid.derive_columns(_hourly_fixture())
    assert out["vre_mw"].iloc[0] == 8000.0        # wind only, solar clipped to 0
    assert out["vre_mw"].iloc[3] == 20000.0       # 8000 wind + 12000 solar


def test_derive_residual_load():
    out = grid.derive_columns(_hourly_fixture())
    expected = out["demand_mw"] - out["vre_mw"]
    pd.testing.assert_series_equal(out["residual_load_mw"], expected, check_names=False)
    assert out["residual_load_mw"].iloc[3] == 175000.0 - 20000.0


def test_derive_generation_total():
    out = grid.derive_columns(_hourly_fixture())
    # thermal + hydro + gas + nuclear + wind + clipped solar
    assert out["generation_mw"].iloc[0] == 100000 + 20000 + 4000 + 5000 + 8000 + 0
    assert out["generation_mw"].iloc[3] == 100000 + 20000 + 4000 + 5000 + 8000 + 12000


def test_derive_is_pure():
    """derive_columns must not mutate its input."""
    frame = _hourly_fixture()
    before = frame.copy()
    grid.derive_columns(frame)
    pd.testing.assert_frame_equal(frame, before)


# --------------------------------------------------------------------------
# Multi-workbook assembly
# --------------------------------------------------------------------------


def test_build_hourly_reindexes_gaps(tmp_path):
    """A missing hour between workbooks becomes an explicit NaN row."""
    import openpyxl

    def write(name, start_hour, hours):
        workbook = openpyxl.Workbook()
        cfg = workbook.active
        cfg.title = "_osi_config"
        cfg["A2"] = json.dumps({"r": 300.0})
        sheet = workbook.create_sheet("Sheet1")
        sheet.append(["Time"] + [""] * 14)
        sheet.append([None] * 15)
        base = pd.Timestamp(f"2022-05-01 {start_hour:02d}:00:00")
        for i in range(hours * 12):
            timestamp = (base + i * pd.Timedelta(minutes=5)).to_pydatetime()
            sheet.append([timestamp, 1.0, 2.0, 3.0, 4.0, 5.0, 6.0,
                          "=f", "=f", "=f", 150000.0, 0, 0, 0, 0])
        workbook.save(tmp_path / name)

    write("May 2022.xlsx", 0, 2)   # hours 00, 01
    write("June 2022.xlsx", 5, 2)  # hours 05, 06 -- gap at 02, 03, 04

    frame, infos = grid.build_hourly(tmp_path)

    assert len(infos) == 2
    assert len(frame) == 7  # 00..06 inclusive, gap filled with NaN rows
    assert frame["demand_mw"].isna().sum() == 3
    assert (frame["n_samples"] == 0).sum() == 3


def test_build_hourly_no_workbooks_raises(tmp_path):
    with pytest.raises(FileNotFoundError, match="No workbooks"):
        grid.build_hourly(tmp_path)
